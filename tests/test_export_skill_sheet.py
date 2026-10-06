import importlib.util
import re
import tempfile
import shutil
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("exporter", ROOT / "scripts/export_skill_sheet.py")
exporter = importlib.util.module_from_spec(spec)
spec.loader.exec_module(exporter)


class ExportTest(unittest.TestCase):
    def test_all_source_lines_and_project_links_are_preserved(self):
        wb = load_workbook(BytesIO(exporter.workbook_bytes()))
        self.assertEqual(wb.sheetnames, ["スキルシート（エンジニア）", "スキル一覧", "案件詳細"])
        values = [str(c.value) for ws in wb for row in ws for c in row if c.value is not None]
        all_text = "\n".join(values)
        _, sections, details = exporter.sources(ROOT)
        for line in sections[exporter.SECTIONS[0]].splitlines():
            if line.strip().startswith("|"):
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                if cells[0] != "項目名" and not re.fullmatch(r":?-+:?", cells[0]):
                    self.assertIn(cells[1], all_text)
        texts = [sections[name] for name in exporter.SECTIONS]
        texts += [p.read_text(encoding="utf-8") for p in details]
        for text in texts:
            for line in text.splitlines():
                line = line.strip()
                if not line or line.startswith("|"):
                    continue
                line = re.sub(r"^#+\s*|^-\s*", "", line)
                line = exporter.plain(line)
                # 案件のラベルは表見出しへ集約し、値が保持されることを確認する。
                if re.match(r"^[^:：()（）]+[:：]", line):
                    line = re.split(r"[:：]", line, maxsplit=1)[1]
                # 項目ラベルと値のセル分割を許容して、本文の欠落を検出する。
                normalized = re.sub(r"[\s:：]", "", line)
                self.assertIn(normalized, re.sub(r"[\s:：]", "", all_text))
        readme = (ROOT / "README.md").read_text(encoding="utf-8")
        links = [url for _, url in exporter.LINK.findall(readme) if "/project/" in url]
        actual = [c.hyperlink.target for row in wb.active for c in row if c.hyperlink]
        self.assertEqual(actual, links)
        for path in details:
            self.assertIn(path.relative_to(ROOT).as_posix(), values)
        for ws in wb:
            self.assertTrue(ws.print_area)
            self.assertEqual(ws.page_setup.fitToWidth, 1)
        profile = wb.active
        self.assertEqual(profile["D3"].value, "ヒラタ トモアキ")
        self.assertEqual(profile["D4"].value, "平田 智昭")
        self.assertIn("東海道線 平塚駅", profile["D5"].value)
        self.assertEqual(profile["I5"].value, "38歳")
        self.assertEqual(profile["I3"].value, "14年目")
        self.assertEqual(profile.page_setup.orientation, "landscape")
        self.assertEqual(profile.freeze_panes, "F14")
        self.assertIn("D3:F3", list(map(str, profile.merged_cells.ranges)))
        self.assertTrue(profile.row_breaks.brk)
        self.assertTrue(wb["スキル一覧"].auto_filter.ref)
        self.assertNotIn("テクフリ 太郎", all_text)

    def test_same_sources_are_reproducible(self):
        self.assertEqual(exporter.workbook_bytes(), exporter.workbook_bytes())

    def test_markdown_edits_change_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "project").mkdir()
            shutil.copytree(ROOT / "excel/templates", root / "excel/templates")
            readme = (ROOT / "README.md").read_text(encoding="utf-8")
            (root / "README.md").write_text(readme, encoding="utf-8")
            detail = root / "project/test.md"
            detail.write_text("### 案件内容\n- 元の経歴\n", encoding="utf-8")
            original = exporter.workbook_bytes(root)
            detail.write_text("### 案件内容\n- 更新した経歴\n", encoding="utf-8")
            self.assertNotEqual(original, exporter.workbook_bytes(root))
            detail_original = exporter.workbook_bytes(root)
            (root / "README.md").write_text(readme.replace("38歳", "39歳"), encoding="utf-8")
            self.assertNotEqual(detail_original, exporter.workbook_bytes(root))

    def test_formula_like_text_is_not_executed(self):
        wb = Workbook()
        exporter.render(wb.active, "- =1+1\n  - 担当：=SUM(A1:A2)\n")
        buffer = BytesIO()
        wb.save(buffer)
        loaded = load_workbook(BytesIO(buffer.getvalue()))
        self.assertEqual(loaded.active["A1"].value, "=1+1")
        self.assertEqual(loaded.active["A1"].data_type, "s")
        self.assertEqual(loaded.active["B2"].data_type, "s")

    def test_completed_side_job_and_explicit_phase_marks(self):
        ws = load_workbook(BytesIO(exporter.workbook_bytes())).active
        # 5月版にあったCRMの「現在」を引き継がず、最新の終了日と副業区分を使用する。
        self.assertEqual(ws["B44"].value, "以降は副業案件です")
        self.assertIn("CRM", ws["F45"].value)
        self.assertEqual(ws["E45"].value, "2026/06")
        self.assertEqual(ws["C47"].value, "11ヶ月")
        # 現案件の「設計」だけで基本設計・詳細設計に丸を付けない。
        self.assertIsNone(ws["M14"].value)
        self.assertIsNone(ws["N14"].value)
        self.assertEqual([ws[f"{c}14"].value for c in "OPQ"], ["●", "●", "●"])
        template = load_workbook(ROOT / "excel/templates/skill-sheet-template.xlsx").active
        self.assertIsNone(template["D4"].value)
        self.assertIsNone(template["F14"].value)
        self.assertTrue(all(c.data_type != "f" for row in ws for c in row))


if __name__ == "__main__":
    unittest.main()
