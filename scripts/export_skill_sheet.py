#!/usr/bin/env python3
"""Markdown のスキルシートを再現可能な .xlsx に出力する。"""

import argparse
import hashlib
import math
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.page import PageMargins
from openpyxl.cell.cell import MergedCell
from copy import copy

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "excel/skill-sheet.xlsx"
SECTIONS = ["【プロフィール】", "【職務要約/自己PR】", "【スキル経験】",
            "【現案件】", "【過去案件: フルタイム案件】", "【過去案件: 副業案件】"]
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
FIXED_TIME = datetime(2000, 1, 1)


def plain(text):
    return LINK.sub(r"\1", text).replace("**", "").replace("`", "")


def append_row(ws, label, value=None, *, heading=False, indent=0, url=None):
    # Excel のセル上限を超える場合は切り捨てずエラーにする。
    for text in (label, value):
        if text is not None and len(text) > 32767:
            raise ValueError("Excel のセル文字数上限を超えています")
    ws.append([label, value])
    row = ws.max_row
    for cell in ws[row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"  # '=' などで始まる経歴も数式として実行しない
        cell.font = Font(name="Yu Gothic", size=11, color="243746")
        cell.alignment = Alignment(vertical="top", wrap_text=True, indent=indent)
        cell.border = Border(bottom=Side(style="hair", color="D9E2EC"))
    if value is None:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    if heading:
        for cell in ws[row]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(name="Yu Gothic", size=12, bold=True, color="FFFFFF")
        ws.row_dimensions[row].height = 32
    else:
        # 日本語の全角幅を考慮し、長文は複数行に分けて描画する。
        widths = (115, 88) if value is None else (25, 88)
        lines = max(sum(max(1, math.ceil(sum(2 if ord(c) > 127 else 1 for c in line) / width))
                        for line in str(text or "").split("\n"))
                    for text, width in zip((label, value), widths))
        ws.row_dimensions[row].height = min(409, max(24, lines * 17 + 8))
    if url:
        ws.cell(row, 1).hyperlink = url


def render(ws, text):
    """既存の見出し、2列表、箇条書き、段落を欠落させず出力する。"""
    paragraph = []

    def flush():
        if paragraph:
            append_row(ws, plain("\n".join(paragraph)))
            paragraph.clear()

    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            flush()
        elif stripped.startswith("#"):
            flush()
            append_row(ws, plain(stripped.lstrip("#").strip()), heading=True)
        elif stripped.startswith("|"):
            flush()
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if all(re.fullmatch(r":?-+:?", c) for c in cells):
                continue
            if len(cells) != 2:
                raise ValueError("プロフィール表は2列で記述してください")
            append_row(ws, plain(cells[0]), plain(cells[1]))
        elif stripped.startswith("- "):
            flush()
            body = stripped[2:]
            indent = min(10, (len(line) - len(line.lstrip())) // 2)
            link = LINK.search(body)
            label = plain(body)
            # 案件名の期間内の ':' は項目区切りにしない。
            field = re.match(r"^([^:：()（）]+)[:：]\s*(.*)$", label)
            if field and indent:
                append_row(ws, field[1], field[2], indent=max(0, indent - 1))
            else:
                append_row(ws, label, indent=indent, url=link[2] if link else None)
        else:
            paragraph.append(stripped)
    flush()


def sources(root):
    readme = (root / "README.md").read_text(encoding="utf-8")
    parts = re.split(r"^## (.+)\s*$", readme, flags=re.MULTILINE)
    sections = dict(zip(parts[1::2], parts[2::2]))
    missing = set(SECTIONS) - sections.keys()
    if missing:
        raise ValueError(f"README の必須セクションがありません: {sorted(missing)}")
    details = sorted((root / "project").rglob("*.md"), reverse=True)
    if not details:
        raise ValueError("project/ の案件詳細がありません")
    return readme.splitlines()[0], sections, details


def styled_cell(ws, row, column, value, *, header=False, shade=False):
    if isinstance(value, str) and len(value) > 32767:
        raise ValueError("Excel のセル文字数上限を超えています")
    cell = ws.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = "s"
    cell.font = Font(name="Yu Gothic", size=11, bold=header,
                     color="FFFFFF" if header else "243746")
    cell.fill = PatternFill("solid", fgColor="17365D" if header else "EDF3F8" if shade else "FFFFFF")
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    cell.border = Border(*(Side(style="thin", color="CED9E5") for _ in range(4)))
    return cell


def merged_row(ws, text, width, *, header=False, height=26):
    if height > 409:
        raise ValueError("Excel の行高上限を超えています。長文を複数行に分割してください")
    row = ws.max_row + 1
    styled_cell(ws, row, 1, text, header=header)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    ws.row_dimensions[row].height = height
    return row


def wrapped_lines(text, width):
    return sum(max(1, math.ceil(sum(2 if ord(c) > 127 else 1 for c in line) / width))
               for line in str(text or "").split("\n"))


def projects(text):
    result = []
    for line in text.splitlines():
        if line.startswith("- "):
            body = line[2:]
            link = LINK.search(body)
            result.append({"title": plain(body), "url": link[2] if link else None, "fields": []})
        elif line.strip().startswith("- ") and result:
            body = plain(line.strip()[2:])
            field = re.match(r"^([^:：]+)[:：]\s*(.*)$", body)
            if not field:
                raise ValueError(f"未対応の案件項目: {body}")
            result[-1]["fields"].append((field[1].strip(), field[2]))
    return result


def engineer_sheet(wb, updated, sections):
    """指定された2026年5月版のセル配置・結合を保って最新情報を転記する。"""
    ws = wb.active
    ws.title = "スキルシート（エンジニア）"
    profile = {}
    for line in sections[SECTIONS[0]].splitlines():
        if not line.strip().startswith("|"):
            continue
        cells = [plain(c.strip()) for c in line.strip().strip("|").split("|")]
        if len(cells) != 2:
            raise ValueError("プロフィール表は2列で記述してください")
        if cells[0] != "項目名" and not re.fullmatch(r":?-+:?", cells[0]):
            profile[cells[0]] = cells[1]
    placements = {"D3": "フリガナ", "D4": "氏名", "I3": "キャリア年数",
                  "I4": "性別", "I5": "年齢", "I6": "学歴"}
    for cell, key in placements.items():
        ws[cell] = profile.get(key, "")
    ws["D5"] = " / ".join(filter(None, (profile.get("最寄駅"), profile.get("稼働希望"))))
    ws.merge_cells("M1:S1")
    ws["M1"] = updated
    ws["M1"].font = Font(name="メイリオ", size=11)
    ws["M1"].alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[1].height = 24
    parts = re.split(r"^### (.+)\s*$", sections[SECTIONS[2]], flags=re.MULTILINE)
    categories = dict(zip(parts[1::2], parts[2::2]))
    ws["D6"] = "\n".join(plain(line.strip()[2:]) for line in categories.get("業務資格", "").splitlines()
                           if line.strip().startswith("- "))
    ws["D8"] = plain(sections[SECTIONS[1]].strip())
    # 新たな自己評価は作らず、現在案件の担当・技術とプロフィールを要点として表示。
    current = projects(sections[SECTIONS[3]])
    fields = dict(current[0]["fields"]) if current else {}
    strengths = ["担当：" + fields.get("担当", "記載なし"),
                 "言語：" + fields.get("言語", "記載なし"),
                 "FW/ライブラリ：" + fields.get("FW/ライブラリ", "記載なし"),
                 "インフラ：" + fields.get("インフラ", "記載なし")]
    strengths += [key + "：" + value for key, value in profile.items()
                  if key not in set(placements.values()) | {"最寄駅", "稼働希望"}]
    ws["D9"] = "\n".join(strengths)
    skill_lines = []
    for category, content in categories.items():
        if category in ("キャリア", "業務資格"):
            continue
        items = [plain(line.strip()[2:]) for line in content.splitlines() if line.strip().startswith("- ")]
        skill_lines.append("・" + category + "：" + "、".join(items))
    ws["D10"] = "\n".join(skill_lines)
    ws["B9"] = "得意分野・現案件"
    for row, cell in [(6, "D6"), (8, "D8"), (9, "D9"), (10, "D10")]:
        width = 82 if row == 6 else 198
        height = wrapped_lines(ws[cell].value, width) * 14 + 12
        if height > 409:
            raise ValueError(f"概要の表示量が行高上限を超えています: {cell}")
        ws.row_dimensions[row].height = max(32, height)
    ws.row_dimensions[10].height = max(180, ws.row_dimensions[10].height + 35)
    for row in (3, 4, 5):
        ws.row_dimensions[row].height = 32
    ws.row_dimensions[13].height = 92
    for col in range(12, 20):
        ws.cell(13, col).alignment = Alignment(horizontal="center", vertical="center", textRotation=90)
    # 空テンプレートの1案件分のスタイルを取り、案件数に応じて同じ構造を複製する。
    block_cells = [(r - 14, c.column, c.value, copy(c._style))
                   for r in range(14, 17) for c in ws[r] if not isinstance(c, MergedCell)]
    merges = [(m.min_row - 14, m.max_row - 14, m.min_col, m.max_col)
              for m in ws.merged_cells.ranges if m.min_row >= 14]
    main = current + projects(sections[SECTIONS[4]])
    side = projects(sections[SECTIONS[5]])
    row = 14
    # 印刷時に概要と経歴を分け、各ページで表の見出しを繰り返す。
    ws.row_breaks.append(Break(id=11))
    page_height = 0
    for index, entry in enumerate(main + side, 1):
        if index == len(main) + 1:
            for col in range(2, 20):
                styled_cell(ws, row, col, None, header=True)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=19)
            ws.cell(row, 2, "以降は副業案件です")
            ws.row_dimensions[row].height = 30
            ws.row_breaks.append(Break(id=row - 1))
            row += 1
            page_height = 30
        if row != 14:
            for offset, col, value, style in block_cells:
                cell = ws.cell(row + offset, col, value)
                cell._style = copy(style)
            for first, last, start, end in merges:
                ws.merge_cells(start_row=row + first, end_row=row + last,
                               start_column=start, end_column=end)
        fields = dict(entry["fields"])
        title = entry["title"]
        period = re.search(r"[（(](\d{4}年.+?)[）)]", title)
        dates = re.findall(r"(\d{4})年(\d+)月", period[1] if period else "")
        start = "/".join((dates[0][0], dates[0][1].zfill(2))) if dates else ""
        end = "/".join((dates[1][0], dates[1][1].zfill(2))) if len(dates) > 1 else "現在" if dates else ""
        duration = re.search(r"[：:]([０-９0-9]+ヶ月)", period[1] if period else "")
        project_name = title[:period.start()] if period else title
        # タイトル・期間の原文も本文に保持し、推測で月数を計算しない。
        narrative = "≪プロジェクト内容≫\n" + title
        narrative += "\n\n≪担当業務≫\n" + fields.get("工程/作業", "記載なし")
        if fields.get("開発手法"):
            narrative += "\n\n≪開発手法≫\n" + fields["開発手法"]
        consumed = {"担当", "言語", "DB", "インフラ", "工程/作業", "開発手法", "チーム体制"}
        tools = "\n".join(key + "：" + value for key, value in entry["fields"] if key not in consumed)
        values = {"B": index, "C": start, "D": "−", "E": end,
                  "F": "■" + project_name, "G": fields.get("担当", ""),
                  "H": fields.get("言語", ""), "I": fields.get("DB", ""),
                  "J": fields.get("インフラ", "") or "—", "K": tools}
        for col, value in values.items():
            ws[f"{col}{row}"] = value
        ws[f"F{row+1}"] = narrative
        ws[f"G{row+1}"] = "チーム構成\n" + fields.get("チーム体制", "記載なし")
        ws[f"C{row+2}"] = duration[1] if duration else "随時更新" if end == "現在" else ""
        phase_text = fields.get("工程/作業", "")
        phase_tokens = re.split(r"[、,]", phase_text)
        phase_checks = ["要件定義" in phase_text, "基本設計" in phase_text,
                        "詳細設計" in phase_text, any(t.strip() in ("開発", "実装") for t in phase_tokens),
                        "単体テスト" in phase_text, "結合テスト" in phase_text,
                        "総合テスト" in phase_text, "保守運用" in phase_text or "保守・運用" in phase_text]
        for col, experienced in enumerate(phase_checks, 12):
            cell = ws.cell(row, col, "●" if experienced else None)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        title_height = max(36, wrapped_lines(project_name, 70) * 14 + 10,
                           wrapped_lines(fields.get("担当", ""), 16) * 14 + 10)
        technical_height = max(wrapped_lines(values[col], width) for col, width in
                               [("H", 19), ("I", 17), ("J", 8), ("K", 18)]) * 14 + 10
        body_height = max(wrapped_lines(narrative, 70) * 14 + 10,
                          wrapped_lines(ws[f"G{row+1}"].value, 16) * 14 + 10,
                          technical_height - title_height - 24)
        if body_height > 409:
            raise ValueError(f"案件の行高上限を超えています: {title}")
        total = title_height + body_height + 24
        if page_height and page_height + total > 590:
            ws.row_breaks.append(Break(id=row - 1))
            page_height = 0
        ws.row_dimensions[row].height = title_height
        ws.row_dimensions[row+1].height = body_height
        ws.row_dimensions[row+2].height = 24
        for c in (f"F{row}", f"G{row}"):
            ws[c].fill = PatternFill("solid", fgColor="EDF3F8")
            ws[c].font = Font(name="メイリオ", size=11, bold=True, color="17365D")
        if entry["url"]:
            ws[f"F{row}"].hyperlink = entry["url"]
        page_height += total
        row += 3
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=19)
    ws.cell(row, 2, "●：Markdownに明記された担当工程。空欄：未記載（経験なしを意味しません）。")
    ws.row_dimensions[row].height = 26
    # 全セルを文字列として保持し、結合セルを含めて書式を整える。
    for cells in ws:
        for cell in cells:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if not isinstance(cell, MergedCell) and cell.row != 13 and not (cell.row >= 14 and cell.column >= 12):
                original = cell.alignment
                cell.alignment = Alignment(horizontal=original.horizontal or "left",
                                           vertical="top", wrap_text=True,
                                           textRotation=original.textRotation or 0)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.print_title_rows = "12:13"
    ws.print_area = f"B1:S{row}"
    ws.freeze_panes = "F14"
    return ws


def skills(wb, updated, text):
    ws = wb.create_sheet("スキル一覧")
    for column, width in zip("ABC", (32, 45, 18)):
        ws.column_dimensions[column].width = width
    for col, value in enumerate(("分類", "技術・サービス", "経験年数"), 1):
        styled_cell(ws, 1, col, value, header=True)
    category = ""
    for line in text.splitlines():
        if line.startswith("### "):
            category = line[4:]
        elif line.strip().startswith("- "):
            value = plain(line.strip()[2:])
            experience = re.match(r"^(.+?)\s+(\d+年(?:未満)?)$", value)
            row = ws.max_row + 1
            values = (category, experience[1] if experience else value, experience[2] if experience else "")
            for col, value in enumerate(values, 1):
                styled_cell(ws, row, col, value, shade=row % 2 == 0)
            ws.row_dimensions[row].height = max(25, max(wrapped_lines(v, w - 2) for v, w in zip(values, (32, 45, 18))) * 17 + 8)
    ws.auto_filter.ref = f"A1:C{ws.max_row}"
    ws.freeze_panes = "B2"
    return ws


def workbook_bytes(root=ROOT):
    updated, sections, details = sources(root)
    wb = load_workbook(root / "excel/templates/skill-sheet-template.xlsx")
    wb.properties.creator = "profile"
    wb.properties.title = "スキルシート"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME

    engineer_sheet(wb, updated, sections)
    skills(wb, updated, sections[SECTIONS[2]])

    ws = wb.create_sheet("案件詳細")
    append_row(ws, "案件詳細 — " + updated, heading=True)
    for path in details:
        if ws.max_row > 1:
            ws.row_breaks.append(Break(id=ws.max_row))
        append_row(ws, path.relative_to(root).as_posix(), heading=True)
        render(ws, path.read_text(encoding="utf-8"))

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 95
    for ws in wb:
        if ws.title == "スキルシート（エンジニア）":
            ws.sheet_view.showGridLines = False
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4, header=0.15, footer=0.15)
            ws.oddFooter.center.text = "&P / &N"
            continue
        ws.freeze_panes = "B2"
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:1"
        ws.print_area = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4, header=0.15, footer=0.15)
        ws.oddFooter.center.text = "&P / &N"

    raw = BytesIO()
    wb.save(raw)
    # ZIP の時刻と openpyxl が更新する modified を固定し、再生成の差分を安定させる。
    result = BytesIO()
    with ZipFile(raw) as original, ZipFile(result, "w", ZIP_DEFLATED) as normalized:
        for name in sorted(original.namelist()):
            data = original.read(name)
            if name == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]+", rb"\g<1>2000-01-01T00:00:00Z", data)
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            normalized.writestr(info, data)
    return result.getvalue()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--check", action="store_true", help="保存済み Excel が最新か検証する")
    args = parser.parse_args()
    data = workbook_bytes()
    if args.check:
        if not args.output.is_file() or args.output.read_bytes() != data:
            parser.exit(1, "Excel が未生成または古い状態です。python scripts/export_skill_sheet.py を実行してください。\n")
        print("Markdown と Excel は同期しています。")
    else:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_bytes(data)
        print(f"出力: {args.output} (SHA256: {hashlib.sha256(data).hexdigest()})")


if __name__ == "__main__":
    main()
